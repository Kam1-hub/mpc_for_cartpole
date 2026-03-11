

import os
from datetime import datetime
import numpy as np
from openpyxl import Workbook


def save_log(log_data):
    """Save simulation results to an XLSX file. Returns the filepath."""

    tracking_mode = log_data.get("tracking_mode", "Static")
    folder_name = f"mpc_{tracking_mode.lower()}"
    base_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), "logs", folder_name)
    os.makedirs(base_dir, exist_ok=True)


    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    filepath = os.path.join(base_dir, f"{timestamp}.xlsx")

    wb = Workbook()


    ws_params = wb.active
    ws_params.title = "Parameters"

    params_data = [
        ("Category", "Parameter", "Value"),
        ("Simulation", "T_FINAL (s)", log_data.get("t_final", "")),
        ("Simulation", "DT (s)", log_data.get("dt", "")),
        ("Initial State", "x0 (m)", log_data["x0"][0]),
        ("Initial State", "q0 (rad)", log_data["x0"][1]),
        ("Initial State", "dx0 (m/s)", log_data["x0"][2]),
        ("Initial State", "dq0 (rad/s)", log_data["x0"][3]),
        ("Tracking", "Mode", log_data.get("tracking_mode", "")),
        ("Tracking", "Target Value", log_data.get("tracking_val", "")),
        ("Q diagonal", "Q11 (x)", log_data["q_diag"][0]),
        ("Q diagonal", "Q22 (q)", log_data["q_diag"][1]),
        ("Q diagonal", "Q33 (dx)", log_data["q_diag"][2]),
        ("Q diagonal", "Q44 (dq)", log_data["q_diag"][3]),
        ("R", "R11", log_data["r_val"]),
        ("S diagonal", "S11 (x)", log_data["s_diag"][0]),
        ("S diagonal", "S22 (q)", log_data["s_diag"][1]),
        ("S diagonal", "S33 (dx)", log_data["s_diag"][2]),
        ("S diagonal", "S44 (dq)", log_data["s_diag"][3]),
        ("MPC", "Np", log_data.get("Np", "")),
        ("MPC", "u_max (N)", log_data.get("u_max", "")),
        ("MPC", "x_max (m)", log_data.get("x_max", "")),
        ("MPC", "q_max (rad)", log_data.get("q_max", "")),
        ("MPC", "xdot_max (m/s)", log_data.get("xdot_max", "")),
        ("MPC", "qdot_max (rad/s)", log_data.get("qdot_max", "")),
        ("MPC", "du_max (N/step)", log_data.get("du_max", "")),
        ("MPC", "rhos_slack", log_data.get("rho_slack", "")),
        ("UI/Phase 3", "Noise Mean (N)", log_data.get("noise_mean", "")),
        ("UI/Phase 3", "Noise Peak (N)", log_data.get("noise_peak", "")),
        ("UI/Phase 3", "Random Seed", log_data.get("random_seed", "")),
        ("UI/Phase 3", "Ghost Trajectory", log_data.get("ghost_trajectory", "")),
    ]
    

    for i, d in enumerate(log_data.get("disturbances", [])):
        params_data.append(("Disturbance", f"Event {i+1}",
                            f"{d['t_start']:.2f}s – {d['t_end']:.2f}s : {d['force']:+.2f} N"))

    for row in params_data:
        ws_params.append(row)

    ws_params.column_dimensions["A"].width = 16
    ws_params.column_dimensions["B"].width = 20
    ws_params.column_dimensions["C"].width = 40


    ws_metrics = wb.create_sheet("Metrics")

    metrics_data = [
        ("Metric", "Value", "Unit"),
        ("Settle Time", log_data.get("settle_time", "N/A"), "s"),
        ("X Peak Error", log_data.get("x_peak", ""), "m"),
        ("X Overshoot", log_data.get("x_overshoot", ""), "m"),
        ("Q Peak Error", log_data.get("q_peak", ""), "rad"),
        ("Q Overshoot", log_data.get("theta_overshoot", ""), "rad"),
        ("Peak Force (umax)", log_data.get("u_peak", ""), "N"),
        ("Total Energy", log_data.get("energy", ""), ""),
        ("Cost Function (J)", log_data.get("cost", ""), ""),
    ]

    for row in metrics_data:
        ws_metrics.append(row)

    ws_metrics.column_dimensions["A"].width = 22
    ws_metrics.column_dimensions["B"].width = 30
    ws_metrics.column_dimensions["C"].width = 10


    ws_ts = wb.create_sheet("TimeSeries")

    ws_ts.append(["t (s)", "x (m)", "q (rad)", "q (deg)",
                  "dx/dt (m/s)", "dq/dt (rad/s)", "F_total (N)", 
                  "F_mpc (N)", "F_ff (N)",
                  "x_ref (m)", "q_ref (rad)", "xdot_ref (m/s)", "qdot_ref (rad/s)", 
                  "Solve Time (ms)",
                  "F_noise (N)", "F_dist (N)", "F_ext_total (N)"])

    t_history = log_data.get("t_history", [])
    states = log_data.get("states_history", [])
    forces = log_data.get("forces_history", [])
    u_mpc = log_data.get("u_mpc_history", [])
    u_ff = log_data.get("u_ff_history", [])
    solve_times = log_data.get("solve_time_history", [])
    x_d = log_data.get("x_d_history", np.zeros_like(states))
    f_noise_h = log_data.get("f_noise_history", np.zeros(len(t_history)))
    f_dist_h = log_data.get("f_dist_history", np.zeros(len(t_history)))
    f_ext_h = log_data.get("f_ext_history", np.zeros(len(t_history)))
    
    min_len = min(len(t_history), len(states), len(forces), len(x_d))

    for i in range(min_len):
        ws_ts.append([
            round(float(t_history[i]), 4),
            round(float(states[i, 0]), 6),
            round(float(states[i, 1]), 6),
            round(float(np.degrees(states[i, 1])), 4),
            round(float(states[i, 2]), 6),
            round(float(states[i, 3]), 6),
            round(float(forces[i]), 4),
            round(float(u_mpc[i]) if i < len(u_mpc) else 0.0, 4),
            round(float(u_ff[i]) if i < len(u_ff) else 0.0, 4),
            round(float(x_d[i, 0]), 6),
            round(float(x_d[i, 1]), 6),
            round(float(x_d[i, 2]), 6),
            round(float(x_d[i, 3]), 6),
            round(float(solve_times[i] * 1000.0) if i < len(solve_times) else 0.0, 4),
            round(float(f_noise_h[i]) if i < len(f_noise_h) else 0.0, 6),
            round(float(f_dist_h[i]) if i < len(f_dist_h) else 0.0, 6),
            round(float(f_ext_h[i]) if i < len(f_ext_h) else 0.0, 6),
        ])

    for col in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q"]:
        ws_ts.column_dimensions[col].width = 14


    wb.save(filepath)
    return filepath


def save_pdf(log_data, xlsx_filepath):
    """Render simulation curves to a PDF file. Returns the PDF filepath."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    pdf_path = xlsx_filepath.rsplit(".", 1)[0] + ".pdf"


    t = np.asarray(log_data.get("t_history", []))
    states = np.asarray(log_data.get("states_history", []))
    forces = np.asarray(log_data.get("forces_history", []))
    x_d = np.asarray(log_data.get("x_d_history", np.zeros_like(states)))
    f_noise = np.asarray(log_data.get("f_noise_history", np.zeros(len(t))))
    f_dist = np.asarray(log_data.get("f_dist_history", np.zeros(len(t))))

    min_len = min(len(t), len(states), len(forces))
    t = t[:min_len]
    states = states[:min_len]
    forces = forces[:min_len]
    x_d = x_d[:min_len]
    f_noise = f_noise[:min_len]
    f_dist = f_dist[:min_len]


    curve_specs = [
        ("x (m)",        "steelblue",    states[:, 0],            x_d[:, 0]     if len(x_d) else None),
        ("q (deg)",      "firebrick",    np.degrees(states[:, 1]), np.degrees(x_d[:, 1]) if len(x_d) else None),
        ("dx/dt (m/s)",  "seagreen",     states[:, 2],            x_d[:, 2]     if len(x_d) else None),
        ("dq/dt (r/s)",  "darkorange",   states[:, 3],            x_d[:, 3]     if len(x_d) else None),
        ("F (N)",        "mediumpurple", forces,                  None),
        ("F_noise (N)",  "darkcyan",     f_noise,                 None),
        ("F_dist (N)",   "orangered",    f_dist,                  None),
    ]


    fig, axes = plt.subplots(7, 1, figsize=(12, 14), sharex=True)
    fig.suptitle("Cart-Pole MPC Simulation Results", fontsize=14, fontweight="bold")

    for i, (label, color, data, ref) in enumerate(curve_specs):
        ax = axes[i]
        ax.plot(t, data, color=color, lw=1.0, label=label)
        if ref is not None:
            ax.plot(t, ref, color="mediumseagreen", ls="--", alpha=0.8, lw=1.0, label="ref")
        ax.set_ylabel(label, fontsize=9)
        ax.tick_params(labelsize=8)
        ax.grid(True, alpha=0.2)
        ax.axhline(y=0, color="k", alpha=0.2, lw=0.5)
        ax.legend(loc="upper right", fontsize=7, framealpha=0.6)

    axes[-1].set_xlabel("t (s)", fontsize=10)


    settle_time = log_data.get("settle_time")
    if settle_time is not None:
        for ax in axes:
            ax.axvline(x=settle_time, color="gray", ls="--", alpha=0.5, lw=0.8)

    fig.tight_layout(rect=[0, 0, 1, 0.97])

    with PdfPages(pdf_path) as pdf:
        pdf.savefig(fig, dpi=150)
    plt.close(fig)

    return pdf_path
